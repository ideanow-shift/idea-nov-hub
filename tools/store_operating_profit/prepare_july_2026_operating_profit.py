#!/usr/bin/env python3
"""Prepare a fixed, review-only July 2026 Store Operations profit package."""
from __future__ import annotations

import argparse
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
import uuid


SOURCE_SHA256 = "97E5F6E360DFCD00CE6E69F974C1E9F19BC8EFFE8B89684FD290946D8872485D"
SOURCE_BYTE_SIZE = 4_012_486
SOURCE_FILE_NAME = "会計_法人管理正本.xlsx"
SOURCE_SYSTEM = "accounting_store_operating_profit_v1"
SOURCE_SHEET = "13_PL_ENTITY_SUMMARY"
DETAIL_SHEET = "12_PL_ENTITY_MONTHLY"
FISCAL_MONTH = "2026-07"
COMPANY_KEY = "0001"
COMPANY_ID = "e4059116-bdb3-4e13-9763-bbc77bdfe062"
COMPANY_MAPPING_ID = "c9214c45-a309-53ac-a073-8f55464d66fe"
ACTOR_EMPLOYEE_ID = "369d9cd5-f6ba-4e53-9428-f631f0893469"
METRIC_CODE = "OPERATING_PROFIT"
DEFINITION_VERSION = "v1"
EXPECTED_DETAIL_ROWS = 3_382
APPROVAL_GATE_SETTING = "app.store_operations_july_profit_execution_approval"
APPROVAL_GATE_VALUE = "OWNER_APPROVAL_REQUIRED_AFTER_FIXED_SHA_REVIEW"


# These are the already confirmed Production mapping records used by June 2026.
# The July workbook label is deliberately mapped explicitly; no fuzzy matching is used.
DIRECT_STORE_MAPPINGS: Mapping[str, Mapping[str, str]] = {
    "KYARA HALF": {
        "mappingSourceKey": "KYARA HALF",
        "mappingId": "1c1271b7-a93b-5c2a-8592-a9e071da32c2",
        "storeId": "ac20934d-ef15-4363-8c2f-759193c7fcc7",
        "storeName": "KYARA HALF池袋",
    },
    "アネックス": {
        "mappingSourceKey": "BASSA ANNEX店",
        "mappingId": "69363eba-a099-5a75-80fe-3bc34fce54a6",
        "storeId": "2980442d-294c-4aae-a9bb-78f530a3a20a",
        "storeName": "BASSAANNEX店",
    },
    "上石神井": {
        "mappingSourceKey": "BASSA上石神井店",
        "mappingId": "a14822f2-661a-5d6f-a585-25a03b96925d",
        "storeId": "1bcba30a-d063-4cdb-be74-425e250aeb25",
        "storeName": "BASSA上石神井店",
    },
    "下井草": {
        "mappingSourceKey": "BASSA下井草店",
        "mappingId": "bc64d0dc-7e9f-5e67-8136-09c00b33c2cb",
        "storeId": "fec1e181-ca5b-482d-a865-3f488f19128f",
        "storeName": "BASSA下井草店",
    },
    "保谷": {
        "mappingSourceKey": "BASSA保谷店",
        "mappingId": "b882a116-f383-57df-97b7-fa1cce042a49",
        "storeId": "73ee82b5-86fa-42e7-ab03-0e075d218dc3",
        "storeName": "BASSA保谷店",
    },
    "所沢": {
        "mappingSourceKey": "BASSA所沢店",
        "mappingId": "d944170a-46f8-54a8-a4df-9db4321b71ec",
        "storeId": "1285ac70-9181-44db-9443-cbd043ab908b",
        "storeName": "BASSA所沢店",
    },
    "東大和": {
        "mappingSourceKey": "BASSA東大和店",
        "mappingId": "df163a99-8426-59f0-9ce5-61600b4bd8aa",
        "storeId": "e7bab6a5-9a8c-4f46-abde-e839ce5bf5e6",
        "storeName": "BASSA東大和店",
    },
    "江古田": {
        "mappingSourceKey": "BASSA江古田店",
        "mappingId": "07628f9d-bfb2-5db0-bc2e-e95b30dd6cf2",
        "storeId": "4e5526cc-9ec7-42aa-ac60-579b6c438c88",
        "storeName": "BASSA江古田店",
    },
    "池袋": {
        "mappingSourceKey": "BASSA池袋店",
        "mappingId": "60853a08-5e79-587a-842f-42a115d6e3a0",
        "storeId": "36c222de-0554-4265-b177-3b68285cc4a4",
        "storeName": "BASSA池袋店",
    },
    "石神井公園": {
        "mappingSourceKey": "BASSA石神井公園店",
        "mappingId": "634e6e8e-ef7e-5dac-bafc-1872bc62b304",
        "storeId": "ad931406-22de-4ba7-a6eb-4502d5c50a91",
        "storeName": "BASSA石神井公園店",
    },
    "立川": {
        "mappingSourceKey": "BASSA立川店",
        "mappingId": "f798a0d9-4ba2-5010-9232-ca1884afe36b",
        "storeId": "5f66193f-d360-4967-b9c7-a100c8ee5e94",
        "storeName": "BASSA立川店",
    },
    "野方": {
        "mappingSourceKey": "BASSA野方店",
        "mappingId": "9cef2bf0-1a22-5c57-91a7-47b183256f73",
        "storeId": "b898c63f-1cc1-42c5-be4f-916f24f49cb6",
        "storeName": "BASSA野方店",
    },
    "高田馬場": {
        "mappingSourceKey": "BASSA高田馬場店",
        "mappingId": "bbb0d75d-3fd7-51e5-8f5b-8e04654a3eb5",
        "storeId": "887da14c-2c0d-46b3-8953-962c7c8dd590",
        "storeName": "BASSA高田馬場店",
    },
}

EXCLUDED_FC_LABELS = frozenset({"久米川", "国分寺", "新所沢", "東久留米", "花小金井"})
REQUIRED_HEADERS = (
    "年月", "法人No", "部門名", "department_type", "当月_営業利益",
    "tax_basis", "source_file", "verification_status",
)


def fail(code: str) -> None:
    raise ValueError(code)


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest().upper()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalized_amount(value: Any) -> str:
    if value is None or isinstance(value, bool):
        fail("OPERATING_PROFIT_MISSING")
    try:
        amount = Decimal(str(value).replace(",", "")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        fail("OPERATING_PROFIT_INVALID")
    return format(amount, "f")


def _uuid(label: str, package_sha256: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"idea-nov:{package_sha256}:{label}"))


def _header_index(header: Sequence[Any]) -> Mapping[str, int]:
    names = [str(value or "").strip() for value in header]
    if len(names) != len(set(name for name in names if name)):
        fail("DUPLICATE_HEADER")
    index = {name: position for position, name in enumerate(names) if name}
    if any(name not in index for name in REQUIRED_HEADERS):
        fail("HEADER_CONTRACT_MISMATCH")
    return index


def build_package_from_rows(
    rows: Sequence[Sequence[Any]], *, detail_row_count: int = EXPECTED_DETAIL_ROWS
) -> Mapping[str, Any]:
    if not rows:
        fail("SOURCE_ROWS_EMPTY")
    if detail_row_count != EXPECTED_DETAIL_ROWS:
        fail("JULY_DETAIL_ROW_COUNT_MISMATCH")
    index = _header_index(rows[0])
    selected: list[tuple[int, Sequence[Any]]] = []
    for excel_row, row in enumerate(rows[1:], start=2):
        if str(row[index["年月"]] or "").strip() != FISCAL_MONTH:
            continue
        if str(row[index["法人No"]] or "").strip() != COMPANY_KEY:
            continue
        if str(row[index["department_type"]] or "").strip() != "SALON":
            continue
        selected.append((excel_row, row))

    labels = [str(row[index["部門名"]] or "").strip() for _, row in selected]
    expected_labels = set(DIRECT_STORE_MAPPINGS) | set(EXCLUDED_FC_LABELS)
    if len(labels) != len(expected_labels) or set(labels) != expected_labels:
        fail("JULY_SALON_POPULATION_MISMATCH")
    if len(labels) != len(set(labels)):
        fail("JULY_SALON_DUPLICATE")

    direct_rows: list[dict[str, Any]] = []
    excluded_rows: list[dict[str, Any]] = []
    for excel_row, row in selected:
        label = str(row[index["部門名"]] or "").strip()
        amount = normalized_amount(row[index["当月_営業利益"]])
        if str(row[index["tax_basis"]] or "").strip() != "税抜":
            fail("TAX_BASIS_MISMATCH")
        if str(row[index["verification_status"]] or "").strip() != "PASS_CURRENT_ONLY":
            fail("SOURCE_VERIFICATION_STATUS_MISMATCH")
        source_file = str(row[index["source_file"]] or "").strip()
        if not source_file:
            fail("SOURCE_LINEAGE_MISSING")

        if label in EXCLUDED_FC_LABELS:
            if Decimal(amount) != 0:
                fail("EXCLUDED_FC_NONZERO")
            excluded_rows.append({
                "sourceRowNumber": excel_row,
                "sourceLabel": label,
                "amount": amount,
                "reason": "FC_OPERATOR_RECORDED_IN_FORMAL_MASTER",
            })
            continue

        mapping = DIRECT_STORE_MAPPINGS[label]
        raw_payload = {
            "fiscalMonth": FISCAL_MONTH,
            "companyKey": COMPANY_KEY,
            "sourceLabel": label,
            "mappingSourceKey": mapping["mappingSourceKey"],
            "metricCode": METRIC_CODE,
            "amount": amount,
            "taxBasis": "net",
            "verificationStatus": "PASS_CURRENT_ONLY",
            "sourceWorkbook": source_file,
            "sourceSheet": SOURCE_SHEET,
            "sourceRowNumber": excel_row,
        }
        direct_rows.append({
            **mapping,
            "sourceRowNumber": excel_row,
            "sourceLabel": label,
            "amount": amount,
            "rawPayload": raw_payload,
            # dbf_ingest.raw_rows enforces a lowercase hexadecimal digest.
            "rawPayloadSha256": sha256_bytes(canonical_json(raw_payload).encode("utf-8")).lower(),
        })

    direct_rows.sort(key=lambda row: row["storeId"])
    excluded_rows.sort(key=lambda row: row["sourceLabel"])
    if len(direct_rows) != 13 or len(excluded_rows) != 5:
        fail("PACKAGE_COUNT_MISMATCH")

    payload: dict[str, Any] = {
        "contractVersion": "store-operating-profit-2026-07-v1",
        "fiscalMonth": FISCAL_MONTH,
        "metricCode": METRIC_CODE,
        "definitionVersion": DEFINITION_VERSION,
        "confirmationStatus": "confirmed",
        "source": {
            "sha256": SOURCE_SHA256,
            "byteSize": SOURCE_BYTE_SIZE,
            "fileName": SOURCE_FILE_NAME,
            "sourceSystem": SOURCE_SYSTEM,
            "sheetName": SOURCE_SHEET,
            "detailSheetName": DETAIL_SHEET,
            "julyDetailRows": detail_row_count,
        },
        "company": {
            "sourceKey": COMPANY_KEY,
            "companyId": COMPANY_ID,
            "mappingId": COMPANY_MAPPING_ID,
        },
        "rows": direct_rows,
        "excludedFcRows": excluded_rows,
        "planned": {
            "sourceFilesInsert": 1,
            "importBatchesInsert": 1,
            "rawRowsInsert": 13,
            "stagingRowsInsert": 13,
            "importEventsInsert": 4,
            "canonicalFactsInsert": 13,
            "batchStatusUpdates": 2,
            "entityMappingsInsert": 0,
            "ddl": 0,
        },
    }
    package_sha256 = sha256_bytes(canonical_json(payload).encode("utf-8"))
    payload["identity"] = {
        "packageSha256": package_sha256,
        "requestFingerprint": sha256_bytes(
            canonical_json({
                "packageSha256": package_sha256,
                "projectRef": "nkmxevmioczcmnldreyo",
                "target": "production-dbf-promotion",
            }).encode("utf-8")
        ),
    }
    payload["identity"]["sourceFileId"] = _uuid("source-file", package_sha256)
    payload["identity"]["batchId"] = _uuid("batch", package_sha256)
    return payload


def load_workbook_rows(path: Path) -> tuple[list[list[Any]], int]:
    try:
        import openpyxl  # type: ignore
    except ImportError as error:  # pragma: no cover - runtime dependency check
        raise RuntimeError("OPENPYXL_REQUIRED_FOR_FIXED_WORKBOOK_READ") from error
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SOURCE_SHEET not in workbook.sheetnames or DETAIL_SHEET not in workbook.sheetnames:
        fail("SOURCE_SHEET_MISSING")
    summary_rows = [list(row) for row in workbook[SOURCE_SHEET].iter_rows(values_only=True)]
    detail_rows = workbook[DETAIL_SHEET].iter_rows(values_only=True)
    detail_header = [str(value or "").strip() for value in next(detail_rows)]
    try:
        month_index = detail_header.index("年月")
    except ValueError:
        fail("DETAIL_HEADER_CONTRACT_MISMATCH")
    detail_count = sum(1 for row in detail_rows if str(row[month_index] or "").strip() == FISCAL_MONTH)
    return summary_rows, detail_count


def _sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def build_execution_sql(package: Mapping[str, Any]) -> str:
    identity = package["identity"]
    rows = package["rows"]
    package_sha = identity["packageSha256"]
    source_file_id = identity["sourceFileId"]
    batch_id = identity["batchId"]
    request_fingerprint = identity["requestFingerprint"]
    expected_json = canonical_json([
        {"storeId": row["storeId"], "amount": row["amount"]} for row in rows
    ])
    raw_json = canonical_json([
        {
            "sourceRowNumber": row["sourceRowNumber"],
            "payload": row["rawPayload"],
            "payloadSha256": row["rawPayloadSha256"],
        }
        for row in rows
    ])
    stage_json = canonical_json([
        {
            "sourceRowNumber": row["sourceRowNumber"],
            "storeMappingId": row["mappingId"],
            "storeId": row["storeId"],
            "amount": row["amount"],
            "normalizedPayload": {
                "definitionVersion": DEFINITION_VERSION,
                "confirmationStatus": "confirmed",
                "sourceWorkbookSha256": SOURCE_SHA256,
                "sourceSheet": SOURCE_SHEET,
                "sourceRowNumber": row["sourceRowNumber"],
                "sourceLabel": row["sourceLabel"],
            },
        }
        for row in rows
    ])
    mapping_json = canonical_json([
        {
            "mappingId": row["mappingId"],
            "sourceKey": row["mappingSourceKey"],
            "storeId": row["storeId"],
        }
        for row in rows
    ])
    return f"""-- REVIEW-ONLY until a separate fixed-SHA Production execution approval.
-- Project: nkmxevmioczcmnldreyo
-- Source SHA-256: {SOURCE_SHA256}
-- Package SHA-256: {package_sha}
-- Request fingerprint: {request_fingerprint}
-- Planned inserts: source_files=1, import_batches=1, raw_rows=13,
-- staging_rows=13, import_events=4, canonical_facts=13; DDL=0.
do $approval_gate$
begin
  if current_setting('{APPROVAL_GATE_SETTING}', true) is distinct from '{APPROVAL_GATE_VALUE}' then
    raise exception using errcode='42501', message='OWNER_APPROVAL_GATE_REQUIRED';
  end if;
end
$approval_gate$;
begin;
set local statement_timeout = '120s';
set local lock_timeout = '15s';
select pg_advisory_xact_lock(hashtextextended({_sql_literal(request_fingerprint)}, 0));

do $store_profit$
declare
  v_actor constant uuid := '{ACTOR_EMPLOYEE_ID}'::uuid;
  v_source_file_id constant uuid := '{source_file_id}'::uuid;
  v_batch_id constant uuid := '{batch_id}'::uuid;
  v_count integer;
  v_result jsonb;
begin
  if not exists (
    select 1 from public.employees where id=v_actor and is_active
  ) then
    raise exception using errcode='22023', message='ACTOR_NOT_ACTIVE';
  end if;
  if not exists (
    select 1 from public.corporations
    where id='{COMPANY_ID}'::uuid and corporation_no='{COMPANY_KEY}' and is_active
  ) then
    raise exception using errcode='22023', message='COMPANY_MASTER_MISMATCH';
  end if;

  if not exists (
    select 1 from dbf_ingest.metric_definitions
    where metric_code='{METRIC_CODE}' and definition_version='{DEFINITION_VERSION}'
      and value_kind='amount' and is_active
  ) then
    raise exception using errcode='22023', message='OPERATING_PROFIT_DEFINITION_MISSING';
  end if;

  if not exists (
    select 1 from dbf_ingest.entity_mappings
    where id='{COMPANY_MAPPING_ID}'::uuid and source_system='{SOURCE_SYSTEM}'
      and entity_type='company' and source_key='{COMPANY_KEY}'
      and company_id='{COMPANY_ID}'::uuid and status='active'
  ) then
    raise exception using errcode='22023', message='COMPANY_MAPPING_MISMATCH';
  end if;

  select count(*) into v_count
  from jsonb_to_recordset({_sql_literal(mapping_json)}::jsonb)
    as x("mappingId" text, "sourceKey" text, "storeId" text)
  join dbf_ingest.entity_mappings m
    on m.id=x."mappingId"::uuid and m.source_system='{SOURCE_SYSTEM}'
   and m.entity_type='store' and m.source_key=x."sourceKey"
   and m.store_id=x."storeId"::uuid and m.status='active'
  join public.stores s on s.id=x."storeId"::uuid and s.is_active
   and s.corporation_id='{COMPANY_ID}'::uuid;
  if v_count <> 13 then
    raise exception using errcode='22023', message='STORE_MAPPING_OR_MASTER_MISMATCH';
  end if;

  if exists (
    select 1 from public.dbf_store_monthly_metric_facts
    where fiscal_month=date '2026-07-01' and metric_code='{METRIC_CODE}' and is_active
  ) then
    raise exception using errcode='23505', message='JULY_OPERATING_PROFIT_ALREADY_ACTIVE';
  end if;
  if exists (select 1 from dbf_ingest.source_files where id=v_source_file_id)
     or exists (select 1 from dbf_ingest.import_batches where id=v_batch_id)
     or exists (
       select 1 from dbf_ingest.source_files
       where source_system='{SOURCE_SYSTEM}' and sha256=lower('{SOURCE_SHA256}')
         and byte_size={SOURCE_BYTE_SIZE}
     ) then
    raise exception using errcode='23505', message='FIXED_PACKAGE_ALREADY_PRESENT';
  end if;

  insert into dbf_ingest.source_files
    (id,sha256,byte_size,original_file_name,media_type,source_system,
     received_by_employee_id,received_via)
  values
    (v_source_file_id,lower('{SOURCE_SHA256}'),{SOURCE_BYTE_SIZE},
     '{SOURCE_FILE_NAME}','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
     '{SOURCE_SYSTEM}',v_actor,'nov_hub_secure_session');

  insert into dbf_ingest.import_batches
    (id,source_file_id,fact_kind,fiscal_month,source_type,status,revision,
     created_by_employee_id)
  values
    (v_batch_id,v_source_file_id,'store_operating_result',date '2026-07-01',
     'accounting_store_operating_profit_exact_v1','owner_review',1,v_actor);

  with input as (
    select * from jsonb_to_recordset({_sql_literal(raw_json)}::jsonb)
      as x("sourceRowNumber" integer, payload jsonb, "payloadSha256" text)
  )
  insert into dbf_ingest.raw_rows(batch_id,source_row_number,payload,payload_sha256)
  select v_batch_id,"sourceRowNumber",payload,"payloadSha256"
  from input order by "sourceRowNumber";

  with input as (
    select * from jsonb_to_recordset({_sql_literal(stage_json)}::jsonb)
      as x("sourceRowNumber" integer,"storeMappingId" text,"storeId" text,
           amount text,"normalizedPayload" jsonb)
  )
  insert into dbf_ingest.staging_rows(
    batch_id,raw_row_id,company_mapping_id,store_mapping_id,company_id,store_id,
    metric_code,amount,quantity,rate,source_row_category,mapping_status,
    validation_status,normalized_payload
  )
  select v_batch_id,r.id,'{COMPANY_MAPPING_ID}'::uuid,x."storeMappingId"::uuid,
         '{COMPANY_ID}'::uuid,x."storeId"::uuid,'{METRIC_CODE}',x.amount::numeric,
         null,null,'detail','resolved','valid',x."normalizedPayload"
  from input x
  join dbf_ingest.raw_rows r
    on r.batch_id=v_batch_id and r.source_row_number=x."sourceRowNumber";

  insert into dbf_ingest.import_events(
    batch_id,event_type,from_status,to_status,actor_employee_id,summary
  ) values
    (v_batch_id,'SOURCE_PARSED','received','parsed',v_actor,
     jsonb_build_object('rowCount',13,'packageSha256','{package_sha}',
       'requestFingerprint','{request_fingerprint}')),
    (v_batch_id,'VALIDATION_COMPLETED','parsed','owner_review',v_actor,
     jsonb_build_object('rowCount',13,'warningCount',0,'packageSha256','{package_sha}',
       'requestFingerprint','{request_fingerprint}'));

  v_result := public.dbf_import_approve_v1(v_actor,v_batch_id);
  if v_result->>'status' <> 'approved' or (v_result->>'rowCount')::integer <> 13 then
    raise exception using errcode='22023', message='APPROVAL_RESULT_MISMATCH';
  end if;
  v_result := public.dbf_import_promote_v1(v_actor,v_batch_id);
  if v_result->>'status' <> 'promoted' or (v_result->>'rowCount')::integer <> 13 then
    raise exception using errcode='22023', message='PROMOTION_RESULT_MISMATCH';
  end if;

  select count(*) into v_count
  from jsonb_to_recordset({_sql_literal(expected_json)}::jsonb)
    as x("storeId" text, amount text)
  join public.dbf_store_monthly_metric_facts f
    on f.batch_id=v_batch_id and f.fiscal_month=date '2026-07-01'
   and f.company_id='{COMPANY_ID}'::uuid and f.store_id=x."storeId"::uuid
   and f.metric_code='{METRIC_CODE}' and f.amount=x.amount::numeric
   and f.quantity is null and f.rate is null
   and f.definition_version='{DEFINITION_VERSION}' and f.status='confirmed' and f.is_active;
  if v_count <> 13 then
    raise exception using errcode='22023', message='CANONICAL_READBACK_MISMATCH';
  end if;
  if (select count(*) from dbf_ingest.raw_rows where batch_id=v_batch_id) <> 13
     or (select count(*) from dbf_ingest.staging_rows where batch_id=v_batch_id) <> 13
     or (select count(*) from dbf_ingest.import_events where batch_id=v_batch_id) <> 4
     or (select count(*) from public.dbf_store_monthly_metric_facts
         where fiscal_month=date '2026-07-01' and metric_code='{METRIC_CODE}'
           and is_active) <> 13
     or (select status from dbf_ingest.import_batches where id=v_batch_id) <> 'promoted' then
    raise exception using errcode='22023', message='POSTFLIGHT_COUNT_MISMATCH';
  end if;
end
$store_profit$;
commit;
"""


def write_artifacts(source_path: Path, output_dir: Path) -> Mapping[str, Any]:
    if sha256_file(source_path) != SOURCE_SHA256:
        fail("SOURCE_SHA_MISMATCH")
    if source_path.stat().st_size != SOURCE_BYTE_SIZE:
        fail("SOURCE_BYTE_SIZE_MISMATCH")
    rows, detail_count = load_workbook_rows(source_path)
    package = build_package_from_rows(rows, detail_row_count=detail_count)
    output_dir.mkdir(parents=True, exist_ok=True)
    package_path = output_dir / "store-operating-profit-2026-07.package.json"
    sql_path = output_dir / "store-operating-profit-2026-07.execution.sql"
    manifest_path = output_dir / "store-operating-profit-2026-07.manifest.json"
    package_path.write_text(
        json.dumps(package, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    sql_path.write_text(build_execution_sql(package), encoding="utf-8", newline="\n")
    manifest = {
        "status": "PASS",
        "productionExecuted": False,
        "sourceSha256": SOURCE_SHA256,
        "packageSha256": package["identity"]["packageSha256"],
        "requestFingerprint": package["identity"]["requestFingerprint"],
        "executionApprovalGate": {
            "setting": APPROVAL_GATE_SETTING,
            "requiredValue": APPROVAL_GATE_VALUE,
        },
        "artifacts": {
            package_path.name: sha256_file(package_path),
            sql_path.name: sha256_file(sql_path),
        },
        "planned": package["planned"],
        "validation": {
            "julyDetailRows": detail_count,
            "directStoreRows": len(package["rows"]),
            "excludedFcZeroRows": len(package["excludedFcRows"]),
            "fuzzyMappings": 0,
            "existingMappingsReused": 14,
        },
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    manifest = write_artifacts(args.source, args.output_dir)
    print(json.dumps({
        "status": manifest["status"],
        "sourceSha256": manifest["sourceSha256"],
        "packageSha256": manifest["packageSha256"],
        "planned": manifest["planned"],
        "validation": manifest["validation"],
        "productionExecuted": False,
    }, ensure_ascii=False, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
